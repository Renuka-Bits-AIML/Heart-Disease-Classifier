"""Build the ten-page DOCX report from the Markdown report and generated figures."""

from __future__ import annotations

import re
from pathlib import Path

import matplotlib
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

matplotlib.use("Agg")

ROOT = Path(__file__).resolve().parents[1]
REPORT_PATH = ROOT / "reports" / "report.md"
OUTPUT_PATH = ROOT / "reports" / "report.docx"
FIGURES_PATH = ROOT / "reports" / "figures"


def clean_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("`", "")
    return text.strip()


def draw_architecture(destination: Path) -> None:
    figure, axes = plt.subplots(figsize=(10, 4.6))
    axes.set_xlim(0, 1)
    axes.set_ylim(0, 1)
    axes.axis("off")

    rows = [
        [
            ("UCI archive", "#dcefe9"),
            ("Validate and parse", "#e9e3d2"),
            ("EDA profile", "#dcefe9"),
            ("Stratified split", "#e9e3d2"),
        ],
        [
            ("MLflow runs\nand artifacts", "#e9e3d2"),
            ("Five-fold tuning\nand comparison", "#dcefe9"),
            ("Logistic Regression\nand Random Forest", "#e9e3d2"),
            ("Fold-local\npreprocessing", "#dcefe9"),
        ],
        [
            ("Held-out evaluation", "#dcefe9"),
            ("Refit selected\npipeline", "#e9e3d2"),
            ("FastAPI\nhealth and predict", "#dcefe9"),
            ("Prometheus metrics\nand deployment", "#e9e3d2"),
        ],
    ]
    x_positions = [0.03, 0.275, 0.52, 0.765]
    y_positions = [0.73, 0.43, 0.13]
    box_width = 0.205
    box_height = 0.17

    for row_index, row in enumerate(rows):
        for column_index, (label, color) in enumerate(row):
            x = x_positions[column_index]
            y = y_positions[row_index]
            axes.text(
                x + box_width / 2,
                y + box_height / 2,
                label,
                ha="center",
                va="center",
                fontsize=9,
                bbox={
                    "boxstyle": "round,pad=0.5",
                    "facecolor": color,
                    "edgecolor": "#354b45",
                    "linewidth": 1,
                },
            )
            draw_arrow = (
                column_index > 0
                if row_index == 1
                else column_index < len(row) - 1
            )
            if draw_arrow:
                direction = -1 if row_index == 1 else 1
                start_x = x if direction == -1 else x + box_width
                end_x = start_x + direction * 0.035
                axes.annotate(
                    "",
                    xy=(end_x, y + box_height / 2),
                    xytext=(start_x, y + box_height / 2),
                    arrowprops={"arrowstyle": "->", "color": "#354b45"},
                )

    axes.annotate(
        "",
        xy=(0.87, 0.61),
        xytext=(0.87, 0.72),
        arrowprops={"arrowstyle": "->", "color": "#354b45"},
    )
    axes.annotate(
        "",
        xy=(0.13, 0.31),
        xytext=(0.38, 0.42),
        arrowprops={"arrowstyle": "->", "color": "#354b45", "connectionstyle": "arc3,rad=0.1"},
    )
    axes.set_title("Heart disease MLOps workflow", fontsize=15, pad=10)
    figure.tight_layout()
    figure.savefig(destination, dpi=180, bbox_inches="tight")
    plt.close(figure)


def add_page_number(document: Document) -> None:
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            continue
        rows.append(cells)
    if not rows:
        return

    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Light Shading Accent 1"
    table.autofit = True
    for row_index, values in enumerate(rows):
        cells = table.rows[0].cells if row_index == 0 else table.add_row().cells
        for column_index, value in enumerate(values):
            cells[column_index].text = clean_markdown(value)
            cells[column_index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[column_index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(7.5)
                    if row_index == 0:
                        run.bold = True


def add_figure_gallery(document: Document, images: list[tuple[str, Path]]) -> None:
    for offset in range(0, len(images), 2):
        table = document.add_table(rows=1, cols=2)
        table.autofit = False
        for column in range(2):
            image_index = offset + column
            cell = table.rows[0].cells[column]
            cell.width = Inches(3.45)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if image_index >= len(images):
                continue
            caption, image_path = images[image_index]
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(image_path), width=Inches(2.55))
            label = cell.add_paragraph(caption)
            label.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in label.runs:
                run.italic = True
                run.font.size = Pt(8)


def build_document() -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(8.5)
    normal.font.color.rgb = RGBColor(38, 48, 48)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    document.styles["Title"].font.size = Pt(22)
    document.styles["Heading 1"].font.size = Pt(17)
    document.styles["Heading 2"].font.size = Pt(12)
    add_page_number(document)

    architecture_path = FIGURES_PATH / "architecture.png"
    architecture_path.parent.mkdir(parents=True, exist_ok=True)
    draw_architecture(architecture_path)

    lines = REPORT_PATH.read_text(encoding="utf-8").splitlines()
    index = 0
    report_title_added = False
    section_one_seen = False

    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue

        if line.startswith("# ") and not report_title_added:
            paragraph = document.add_paragraph(clean_markdown(line[2:]), style="Title")
            paragraph.paragraph_format.space_after = Pt(10)
            report_title_added = True
            index += 1
            continue

        if line.startswith("## "):
            match = re.match(r"## (\d+)\.", line)
            if match and int(match.group(1)) > 1:
                document.add_page_break()
            elif match and int(match.group(1)) == 1:
                section_one_seen = True
            heading = document.add_paragraph(clean_markdown(line[3:]), style="Heading 1")
            heading.paragraph_format.keep_with_next = True
            index += 1
            continue

        if line.startswith("```mermaid"):
            while index < len(lines) and not lines[index].strip().startswith("```"):
                index += 1
            if index < len(lines):
                index += 1
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run().add_picture(str(architecture_path), width=Inches(6.5))
            continue

        if line.startswith("!["):
            image_group = []
            while index < len(lines):
                candidate = lines[index].strip()
                match = re.match(r"!\[([^\]]+)\]\(([^)]+)\)", candidate)
                if not match:
                    break
                image_group.append(
                    (match.group(1), REPORT_PATH.parent / match.group(2))
                )
                index += 1
                while index < len(lines) and not lines[index].strip():
                    index += 1
            add_figure_gallery(document, image_group)
            continue

        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            add_markdown_table(document, table_lines)
            continue

        paragraph_lines = [line]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", "|", "![", "```"))
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph(clean_markdown(" ".join(paragraph_lines)))
        paragraph.paragraph_format.keep_together = True

    if not section_one_seen:
        raise ValueError("The Markdown source must include section 1")

    document.core_properties.title = "Heart Disease Classification: MLOps Project Report"
    document.core_properties.subject = "Ten-page report with measured model results"
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(f"Wrote {build_document()}")